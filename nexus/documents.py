from io import BytesIO

from docx import Document
from pypdf import PdfReader


def extract_document(uploaded_file, max_chars: int = 24000) -> tuple[str, dict[str, object]]:
    name = getattr(uploaded_file, "name", "uploaded file")
    suffix = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    payload = uploaded_file.getvalue()
    if not payload:
        raise ValueError("The uploaded file is empty.")

    if suffix == "pdf":
        reader = PdfReader(BytesIO(payload))
        if reader.is_encrypted:
            raise ValueError("This PDF is password protected and cannot be read.")
        pages = [page.extract_text() or "" for page in reader.pages]
        text = "\n\n".join(part.strip() for part in pages if part.strip())
    elif suffix == "docx":
        document = Document(BytesIO(payload))
        text = "\n\n".join(p.text.strip() for p in document.paragraphs if p.text.strip())
        tables = []
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    tables.append(" | ".join(cells))
        if tables:
            text = "\n\n".join(part for part in (text, "\n".join(tables)) if part)
    elif suffix == "txt":
        text = payload.decode("utf-8-sig", errors="replace").strip()
    else:
        raise ValueError("Unsupported file type. Upload PDF, DOCX, or TXT.")

    if not text.strip():
        raise ValueError("No readable text was found. Scanned PDFs may need OCR before upload.")

    was_truncated = len(text) > max_chars
    bounded_text = text[:max_chars]
    return bounded_text, {
        "filename": name,
        "characters_extracted": len(text),
        "characters_sent": len(bounded_text),
        "truncated": was_truncated,
    }
